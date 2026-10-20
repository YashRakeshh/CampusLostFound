from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\Lakshay Shrivastava\Documents\Codex\2026-08-27\project-campus-lost-found-platform-a")
TEMPLATE = ROOT / "work" / "template" / "Project template.docx"
OUTPUT = ROOT / "outputs" / "Campus_Lost_and_Found_Platform_Project_Report.docx"
OUTPUT.parent.mkdir(exist_ok=True)

doc = Document(str(TEMPLATE))
source_list_numpr = deepcopy(doc.paragraphs[35]._p.pPr.numPr)
section = doc.sections[0]
section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)

# Keep only the supplied cover, adapting the one supplied title slot.
doc.paragraphs[0].text = "Campus Lost & Found Platform"
for run in doc.paragraphs[0].runs:
    run.font.name = "Times New Roman"; run.font.size = Pt(16); run.bold = True

# Delete supplied contents/instruction area after the cover. Keep the cover through date.
for p in list(doc.paragraphs[32:]):
    p._element.getparent().remove(p._element)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Times New Roman'; normal._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman'); normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
normal.font.size = Pt(12)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15
for n, size in [('Heading 1', 14), ('Heading 2', 12), ('Heading 3', 12)]:
    try: s = styles[n]
    except KeyError: s = styles.add_style(n, WD_STYLE_TYPE.PARAGRAPH)
    s.font.name = 'Times New Roman'; s._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'), 'Times New Roman'); s._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'), 'Times New Roman'); s.font.size = Pt(size); s.font.bold = True
    s.paragraph_format.space_before = Pt(12); s.paragraph_format.space_after = Pt(6)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run('Campus Lost & Found Platform | ')
r.font.name='Times New Roman'; r.font.size=Pt(9)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
footer._p.append(fld)

def para(text='', style=None, bold_lead=None):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        x=p.add_run(bold_lead); x.bold=True; p.add_run(text[len(bold_lead):])
    else: p.add_run(text)
    for run in p.runs:
        run.font.name='Times New Roman'; run.font.size=Pt(12)
    return p

def h(text, level=1): return para(text, f'Heading {level}')
def bullet(text):
    p=doc.add_paragraph(style='List Paragraph'); p._p.get_or_add_pPr().append(deepcopy(source_list_numpr)); p.add_run(text)
    for r in p.runs: r.font.name='Times New Roman'; r.font.size=Pt(12)
    return p
def number(text):
    p=doc.add_paragraph(style='List Paragraph'); p._p.get_or_add_pPr().append(deepcopy(source_list_numpr)); p.add_run(text)
    for r in p.runs: r.font.name='Times New Roman'; r.font.size=Pt(12)
    return p
def page(): doc.add_page_break()
def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
def set_cell(cell, text, header=False):
    cell.text=text; cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.name='Times New Roman'; r.font.size=Pt(10); r.bold=header
    if header: shade(cell, 'D9EAF7')
def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = OxmlElement(f'w:{edge}'); tag.set(qn('w:val'),'single'); tag.set(qn('w:sz'),'4'); tag.set(qn('w:color'),'808080'); borders.append(tag)
    t._tbl.tblPr.append(borders)
    for c,v in zip(t.rows[0].cells,headers): set_cell(c,v,True)
    trPr = t.rows[0]._tr.get_or_add_trPr(); repeat = OxmlElement('w:tblHeader'); repeat.set(qn('w:val'), 'true'); trPr.append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for c,v in zip(cells,row): set_cell(c,str(v))
    if widths:
        for row in t.rows:
            for c,w in zip(row.cells,widths): c.width=Inches(w)
    doc.add_paragraph()
    return t

page()
title=para('PROJECT REPORT', None); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in title.runs: r.bold=True; r.font.size=Pt(16)
sub=para('Campus Lost & Found Platform', None); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in sub.runs: r.bold=True; r.font.size=Pt(14)
para('Database Systems Lab (BCSE302P)', None).alignment=WD_ALIGN_PARAGRAPH.CENTER
para('Fall Semester 2026-2027', None).alignment=WD_ALIGN_PARAGRAPH.CENTER

h('Abstract')
para('Campus Lost & Found Platform is a database-driven web application for helping students report lost and found belongings and recover them safely. Students submit item details such as category, title, description, image, location, date, and optional reward. The system evaluates potential lost-found pairs using a weighted score based on category agreement, location similarity, date proximity, and keyword similarity. Candidate pairs are stored for review rather than being automatically finalized. A claimant can request an item, after which an administrator verifies the evidence and locks the claim through a transaction so that competing claims cannot both succeed. The proposed implementation uses React for the interface, Node.js and Express for the API, and PostgreSQL for relational storage. The design demonstrates advanced database concepts through normalized tables, multi-table joins, weighted scoring, indexing experiments, notifications, role-based authorization, and row-level locking. The outcome is a practical campus service that is more reliable than an informal notice-board process and a focused demonstration of database systems concepts.')

h('Contents')
for item in ['Introduction','Background Study','Requirements and Scope','System Methodology','Database Design and ER Model','Matching Engine and SQL Design','Modules and Submodules','Transaction and Concurrency Control','Indexing and Performance Evaluation','Results and Discussion','Conclusion','References']:
    bullet(item)

page(); h('1. Introduction')
para('Students frequently lose personal items such as ID cards, wallets, chargers, headphones, keys, books, and laboratory equipment. Existing recovery methods rely on informal WhatsApp groups, security desks, and physical notice boards. These channels are difficult to search, do not preserve evidence consistently, and are vulnerable to duplicate or fraudulent claims. A centralized platform can record each report, search related records, and create an auditable handover process.')
h('1.1 Problem Statement',2)
para('The campus needs a reliable way to connect a lost-item report with a corresponding found-item report while ensuring that an item is handed to only the rightful claimant. The platform must rank likely matches, allow staff verification, and remain correct when more than one student tries to claim the same item.')
h('1.2 Objectives',2)
for x in ['Provide separate workflows for lost and found item reporting.', 'Suggest candidate matches through database queries and weighted scoring.', 'Allow claim requests with evidence while retaining an administrator approval step.', 'Prevent duplicate final claims using ACID transactions and row-level locks.', 'Measure how indexes improve lookup and matching query performance.']: bullet(x)
h('1.3 Scope',2)
para('The first release serves authenticated campus users and administrators. It supports structured descriptions and image URLs, not image recognition. The matching score is decision support: final identity verification remains an administrator responsibility.')

page(); h('2. Background Study')
para('Lost-and-found systems are commonly implemented as simple classified listings. Although such systems support posting and browsing, they often lack structured comparison, claim-state control, and an auditable transfer record. In a campus setting, similarity requires several independent signals because item descriptions vary and the same category can contain many near-identical objects.')
h('2.1 Database Concepts Applied',2)
table(['Concept','Use in the Platform'],[
['Normalization','Separates users, reports, matches, claims, and notifications to reduce redundancy.'],
['Joins','Combines lost reports, found reports, reporter details, and claim state for matching and review.'],
['Transactions','Makes claim approval all-or-nothing when concurrent claimants target one found item.'],
['Indexes','Accelerate category/date filters and keyword search used by the matching engine.'],
['Authorization','Restricts final verification and status changes to administrators.']],[1.5,5.0])
h('2.2 Design Assumptions',2)
para('Every report belongs to one registered user. A found item can have many proposed matches and claim requests over time, but only one approved claim. Dates represent the approximate loss or discovery date, and locations are selected from a campus location vocabulary where possible.')

page(); h('3. Requirements and Scope')
h('3.1 Functional Requirements',2)
for x in ['Register and authenticate students and administrators.', 'Create, edit, close, and browse lost/found reports.', 'Attach category, location, date, descriptive keywords, image URL, and optional reward.', 'Generate and rank possible lost-found matches.', 'Submit a claim request with identifying evidence.', 'Allow an administrator to approve or reject a request and notify both parties.', 'Maintain status history through report, match, and claim records.']: number(x)
h('3.2 Non-functional Requirements',2)
for x in ['Correctness of finalized claims under concurrent requests.', 'Fast candidate retrieval for growing datasets.', 'Privacy: only necessary contact details appear to authorized users.', 'Auditability through timestamps and state transitions.', 'Responsive interface suitable for mobile and desktop browsers.']: bullet(x)

h('4. System Methodology')
para('The application follows a three-tier architecture. React presents forms, feeds, match suggestions, and admin review screens. Node.js/Express validates requests, applies authorization, and executes parameterized SQL. PostgreSQL persists transactional data and evaluates match queries. Uploaded images are stored externally; the database records their URLs and metadata.')
h('4.1 Core Workflow',2)
number('Student submits a lost or found item report with structured metadata.')
number('Matching engine filters eligible opposite-type reports and computes a score for each pair.')
number('High-scoring candidates are stored in Matches and shown to the relevant users/admin.')
number('A student submits a claim request with evidence for a found item.')
number('Administrator verifies evidence, approves one claim, locks the item, and triggers notifications.')
h('4.2 Weighted Match Formula',2)
para('score = 0.35 × category_score + 0.25 × location_score + 0.20 × date_score + 0.20 × keyword_score. Each component is normalized to 0–100. Candidate matches above a configurable threshold (for example, 55) are persisted. The weights can be calibrated after collecting feedback from verified and rejected matches.')

page(); h('5. Database Design and ER Model')
para('The schema is normalized to third normal form. Item reports retain their report-specific attributes in LostItems and FoundItems, while shared account and workflow information is represented by Users, Matches, Claims, and Notifications.')
table(['Entity','Primary Key','Purpose / Important Attributes'],[
['Users','user_id','name, email, role, phone, created_at'],
['LostItems','lost_id','reporter_id, category, title, description, location, lost_date, image_url, reward, status'],
['FoundItems','found_id','reporter_id, category, title, description, location, found_date, image_url, status'],
['Matches','match_id','lost_id, found_id, component scores, overall_score, status, created_at'],
['Claims','claim_id','match_id, claimant_id, evidence, status, reviewed_by, reviewed_at'],
['Notifications','notification_id','user_id, type, message, related_claim_id, read_at']
],[1.25,1.2,4.05])
h('5.1 Relationship Summary',2)
para('Users 1:N LostItems and Users 1:N FoundItems; LostItems N:M FoundItems through Matches; Matches 1:N Claims; Users 1:N Claims as claimants; Users 1:N Notifications. A partial unique rule ensures that only one approved claim exists for a given found item (enforced through the match relationship).')
h('5.2 Logical ER Diagram',2)
para('Users ──< LostItems ──< Matches >── FoundItems >── Users\n                 │                  │\n                 └──────────────< Claims >── Users (claimant/admin)\n                                      │\nUsers ────────────────────────────< Notifications')

page(); h('6. Matching Engine and SQL Design')
para('The candidate query first limits records by compatible category and unresolved state. It then calculates location, date, and keyword components. PostgreSQL full-text search or trigram similarity can derive the keyword score; the database can materialize or recompute the final score when a report changes.')
h('6.1 Illustrative Candidate Query',2)
code = "SELECT l.lost_id, f.found_id,\n  (CASE WHEN l.category=f.category THEN 100 ELSE 0 END) AS category_score,\n  GREATEST(0, 100 - 10*ABS(l.lost_date-f.found_date)) AS date_score,\n  similarity(lower(l.description), lower(f.description))*100 AS keyword_score\nFROM lost_items l JOIN found_items f ON l.category=f.category\nWHERE l.status='open' AND f.status='open';"
p=para(code); p.paragraph_format.left_indent=Inches(.25); p.paragraph_format.space_after=Pt(10)
for r in p.runs: r.font.name='Courier New'; r.font.size=Pt(9)
h('6.2 Match Lifecycle',2)
table(['Status','Meaning'],[['suggested','Generated candidate awaiting interest or review.'],['claimed','A claimant has requested possession.'],['verified','Administrator has approved the claimant.'],['rejected','Candidate or claim was rejected.'],['closed','Associated item is no longer available.']],[1.4,5.1])

page(); h('7. Modules and Submodules')
table(['Module','Submodules','Responsibilities'],[
['Authentication','Registration, login, role check','Creates sessions and distinguishes student/admin actions.'],
['Item Reporting','Lost report, found report, image/link validation','Captures data and keeps report status current.'],
['Matching','Candidate query, score calculation, thresholding','Ranks opposite-type records and saves suggested pairs.'],
['Claims','Evidence form, request history, status updates','Lets students request a specific item safely.'],
['Administration','Review queue, approval/rejection, audit view','Verifies identity evidence and finalizes a handover.'],
['Notifications','In-app event creation, read status','Informs users about matches, decisions, and handover actions.']
],[1.25,2.15,3.1])

h('8. Transaction and Concurrency Control')
para('Claim approval is the critical transaction. Two administrators or two claimants may act nearly simultaneously on the same found item. The database must serialize the finalization decision so that only one request becomes verified.')
h('8.1 Approval Transaction',2)
number('Begin a transaction and lock the target FoundItems row with SELECT ... FOR UPDATE.')
number('Recheck that the item is open and that no approved claim already exists.')
number('Lock and update the chosen claim to verified; reject competing pending claims for the same found item.')
number('Mark the found item and matched lost item as closed; insert notifications for both reporters.')
number('Commit. On an error, roll back every update.')
para('The lock prevents a second transaction from seeing the item as available until the first transaction commits or rolls back. This directly demonstrates isolation and atomicity from the DBMS syllabus.')

page(); h('9. Indexing and Performance Evaluation')
para('The matching workload filters category and report status, compares dates, and searches descriptions. Indexes reduce the number of candidate rows examined before the weighted scoring expression runs. Performance should be evaluated using EXPLAIN ANALYZE on representative synthetic data at increasing sizes.')
table(['Index','Purpose'],[
['idx_lost_category_status_date','Composite B-tree on (category, status, lost_date) for active lost-report filtering.'],
['idx_found_category_status_date','Composite B-tree on (category, status, found_date) for active found-report filtering.'],
['idx_lost_keywords / idx_found_keywords','GIN index on generated tsvector description/title data for keyword lookup.'],
['idx_matches_pair','Unique (lost_id, found_id) to prevent duplicate match rows.'],
['idx_claims_pending','Partial index for pending claims used by the review queue.']
],[2.2,4.3])
h('9.1 Benchmark Plan',2)
for x in ['Load 10,000–100,000 reports with realistic category and date distributions.', 'Run the same candidate query before creating indexes and record planning/execution time.', 'Create indexes, ANALYZE the tables, and repeat the query.', 'Compare execution time, rows scanned, and query plan nodes; document write-cost trade-offs.']: number(x)

page(); h('10. Results and Discussion')
para('The expected prototype result is a ranked review queue where obvious pairs appear near the top and ambiguous pairs remain available for human confirmation. The score is explainable because each component can be displayed to the administrator, such as “same category, nearby location, two-day date gap, strong keyword overlap.” This reduces overreliance on an opaque automated decision.')
h('10.1 Validation Scenarios',2)
table(['Scenario','Expected Result'],[
['Strong match','Same category, nearby location, close date, similar description yields a high score.'],
['Weak match','Category differs or date gap is large; record is excluded or ranked low.'],
['Duplicate claim','Second claim remains pending/rejected after first approved transaction locks item.'],
['Admin rejection','Claim is rejected; found item stays open for another legitimate claimant.'],
['Closed item','No new approved claim can be created for the item.']
],[1.6,4.9])
h('10.2 Limitations and Future Work',2)
para('Scores depend on accurate user descriptions and selected locations. Future versions can add image embeddings, campus-map distance, multilingual keyword processing, email/SMS delivery, and analytics for locations with repeated losses. Such additions should retain human verification for privacy and fraud prevention.')

h('11. Conclusion')
para('Campus Lost & Found Platform converts an informal recovery process into a reliable transactional system. Its core contribution is not merely posting records: it combines normalized relational design, scored multi-table matching, index-aware query optimization, and row-locking for safe claim finalization. The design is feasible with React, Node/Express, and PostgreSQL and provides a strong, demonstrable Database Systems Lab project.')

h('12. References')
for x in ['PostgreSQL Global Development Group. PostgreSQL Documentation: Transactions, Explicit Locking, and Indexes.', 'Node.js Foundation. Node.js Documentation.', 'React Team. React Documentation.', 'Elmasri, R. and Navathe, S. Fundamentals of Database Systems. Pearson.', 'Silberschatz, A., Korth, H. F., and Sudarshan, S. Database System Concepts. McGraw-Hill.']:
    bullet(x)

doc.save(str(OUTPUT))
print(OUTPUT)
