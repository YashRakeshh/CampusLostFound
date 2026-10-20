from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(r'C:\Users\Lakshay Shrivastava\Documents\Codex\2026-08-27\project-campus-lost-found-platform-a')
doc=Document(str(root/'work'/'review-template'/'Project review -1 document.docx'))
out=root/'outputs'/'Campus_Lost_and_Found_Platform_Review_1.docx'; out.parent.mkdir(exist_ok=True)
doc.paragraphs[0].text='Campus Lost & Found Platform'
for r in doc.paragraphs[0].runs: r.font.name='Times New Roman'; r.font.size=Pt(16); r.bold=True

# Preserve the supplied cover, and replace only the prompt/instruction section.
for p in list(doc.paragraphs[32:]): p._element.getparent().remove(p._element)
normal=doc.styles['Normal']; normal.font.name='Times New Roman'; normal._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); normal._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman'); normal.font.size=Pt(12)

def p(text='', bold=False, size=12, align=None, before=0, after=6):
    x=doc.add_paragraph(); x.paragraph_format.space_before=Pt(before); x.paragraph_format.space_after=Pt(after)
    if align is not None: x.alignment=align
    r=x.add_run(text); r.font.name='Times New Roman'; r.font.size=Pt(size); r.bold=bold
    return x
def heading(text): return p(text,True,12,None,12,6)
def table(headers, rows):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    b=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'808080'); b.append(el)
    t._tbl.tblPr.append(b)
    for cell,text in zip(t.rows[0].cells,headers):
        cell.text=text
        for run in cell.paragraphs[0].runs: run.font.name='Times New Roman';run.font.size=Pt(10);run.bold=True
        shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'D9EAD3');cell._tc.get_or_add_tcPr().append(shd)
    tr=t.rows[0]._tr.get_or_add_trPr(); repeat=OxmlElement('w:tblHeader'); repeat.set(qn('w:val'),'true');tr.append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for cell,text in zip(cells,row):
            cell.text=text
            for run in cell.paragraphs[0].runs: run.font.name='Times New Roman';run.font.size=Pt(10)
    doc.add_paragraph()

doc.add_page_break()
p('PROJECT REVIEW - 1',True,16,WD_ALIGN_PARAGRAPH.CENTER,0,6)
p('Campus Lost & Found Platform',True,14,WD_ALIGN_PARAGRAPH.CENTER,0,18)
heading('Project Title')
p('Campus Lost & Found Platform: a web application that helps students report lost or found items and suggests possible matches using category, location, date proximity, and keyword similarity.')
heading('Short Summary of the Project')
p('The platform provides a structured alternative to informal lost-and-found groups. A student can create a lost or found report with item details, location, date, image reference, and optional reward. The matching engine compares opposite-type reports and assigns an explainable weighted score. Potential matches are reviewed before a claimant can receive an item. An administrator verifies the claim and the system locks the item so it cannot be finalized twice.')
heading('Current Progress: 25% Complete')
p('The initial foundation has been completed. The project structure, user-facing report flow, database design, score-based matching logic, and a local PostgreSQL development environment are implemented. The next phase will focus on strengthening workflow coverage, adding administrator review screens, and completing end-to-end validation.')
table(['Work Area','Status at Review 1'],[
['Project setup and architecture','Completed'],
['Database schema and seed data','Completed'],
['Lost/found report creation UI','Completed'],
['Weighted match query design','Completed'],
['Claim approval and notification workflow','Planned for next phase'],
['Authentication and full test coverage','Planned for next phase']])

heading('Front End Details')
p('The front end is built with React and Vite. It currently provides a responsive landing view, a report form for lost/found items, a feed of open reports, and a match-results view. The match score is shown with its main components so that users and reviewers can understand how a suggestion was generated. The interface communicates with the backend through REST API calls.')
heading('Database Details')
p('PostgreSQL is used as the relational database. The current schema includes Users, Items, Matches, Claims, and Notifications. The Items table records both lost and found reports using a kind field. The Matches table stores category, location, date, keyword, and overall scores for each candidate pair. Composite and text-search indexes have been planned/created for category, status, date, and description matching.')
heading('Next Steps')
p('The upcoming work will implement the administrator claim-review dashboard, complete row-locking based claim approval, add notifications for match and claim events, test concurrent claim attempts, and document index performance before and after optimization. This will raise the project from the foundational 25% stage to a complete working prototype.')
doc.save(str(out)); print(out)
